from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "SprintFlow_presentation_mn.pptx"
LANDING_IMAGE = ROOT / "frontend" / "public" / "landing-team.jpg"
HERO_IMAGE = ROOT / "frontend" / "src" / "assets" / "hero.png"

BG = RGBColor(10, 18, 33)
PANEL = RGBColor(18, 32, 56)
PANEL_SOFT = RGBColor(235, 242, 255)
ACCENT = RGBColor(59, 130, 246)
ACCENT_2 = RGBColor(124, 58, 237)
TEXT = RGBColor(241, 245, 249)
TEXT_DARK = RGBColor(15, 23, 42)
TEXT_MUTED = RGBColor(148, 163, 184)
LINE = RGBColor(203, 213, 225)
GOOD = RGBColor(16, 185, 129)


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, size=24, bold=False, color=TEXT_DARK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    run = p.runs[0]
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Arial"
    return box


def add_bullets(slide, left, top, width, height, items, size=19, color=TEXT_DARK):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"• {item}"
        p.level = 0
        p.space_after = Pt(10)
        run = p.runs[0]
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = "Arial"
    return box


def add_title(slide, eyebrow, title, subtitle=None, light=False):
    text_color = TEXT if light else TEXT_DARK
    muted = RGBColor(147, 197, 253) if light else ACCENT
    add_textbox(slide, Inches(0.8), Inches(0.55), Inches(5.8), Inches(0.4), eyebrow, size=12, bold=True, color=muted)
    add_textbox(slide, Inches(0.8), Inches(0.95), Inches(6.2), Inches(1.1), title, size=28, bold=True, color=text_color)
    if subtitle:
        add_textbox(slide, Inches(0.8), Inches(1.9), Inches(6.8), Inches(0.7), subtitle, size=16, color=TEXT_MUTED if light else RGBColor(71, 85, 105))


def add_panel(slide, left, top, width, height, fill_rgb, line_rgb=None, radius=0.12):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        left,
        top,
        width,
        height,
    )
    shape.adjustments[0] = radius
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = fill_rgb
    line = shape.line
    line.color.rgb = line_rgb or fill_rgb
    line.width = Pt(1.2)
    return shape


def add_metric(slide, left, top, width, height, value, label, accent):
    add_panel(slide, left, top, width, height, PANEL_SOFT, LINE, radius=0.08)
    add_textbox(slide, left + Inches(0.18), top + Inches(0.15), width - Inches(0.3), Inches(0.45), value, size=22, bold=True, color=accent)
    add_textbox(slide, left + Inches(0.18), top + Inches(0.55), width - Inches(0.3), Inches(0.55), label, size=12, color=TEXT_DARK)


def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, BG)
    slide.shapes.add_picture(str(LANDING_IMAGE), Inches(7.0), Inches(0), Inches(6.33), Inches(7.5))
    overlay = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(6.6), Inches(0), Inches(6.73), Inches(7.5))
    overlay.fill.solid()
    overlay.fill.fore_color.rgb = RGBColor(10, 18, 33)
    overlay.fill.transparency = 0.35
    overlay.line.fill.background()

    add_textbox(slide, Inches(0.8), Inches(0.8), Inches(5.8), Inches(0.5), "ДИПЛОМЫН ТӨСӨЛ", size=13, bold=True, color=RGBColor(147, 197, 253))
    add_textbox(slide, Inches(0.8), Inches(1.35), Inches(5.7), Inches(1.25), "SprintFlow", size=30, bold=True, color=TEXT)
    add_textbox(slide, Inches(0.8), Inches(2.35), Inches(5.8), Inches(1.1), "Agile төслийн менежментийн ухаалаг вэб систем", size=24, bold=True, color=TEXT)
    add_textbox(
        slide,
        Inches(0.8),
        Inches(3.45),
        Inches(5.5),
        Inches(1.0),
        "Төсөл, даалгавар, багийн уялдаа болон AI туслахыг нэг платформд нэгтгэсэн вэб систем.",
        size=16,
        color=RGBColor(203, 213, 225),
    )
    add_panel(slide, Inches(0.8), Inches(5.35), Inches(2.15), Inches(0.72), PANEL, RGBColor(51, 65, 85))
    add_textbox(slide, Inches(1.02), Inches(5.57), Inches(1.8), Inches(0.2), "React + Express + PostgreSQL", size=11, color=TEXT)
    add_textbox(slide, Inches(0.8), Inches(6.55), Inches(3.0), Inches(0.25), "Оюутны танилцуулгын богино хувилбар", size=11, color=TEXT_MUTED)


def slide_problem_goal(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, RGBColor(248, 250, 252))
    add_title(
        slide,
        "АСУУДАЛ БА ЗОРИЛГО",
        "Яагаад SprintFlow хэрэгтэй вэ?",
        "Өдөр тутмын ажлын тархмал мэдээллийг нэг цэгт төвлөрүүлж, багийн хариуцлага ба явцыг илүү ойлгомжтой болгоно.",
    )
    add_panel(slide, Inches(0.8), Inches(2.7), Inches(5.7), Inches(3.7), RGBColor(255, 255, 255), LINE)
    add_bullets(
        slide,
        Inches(1.1),
        Inches(3.05),
        Inches(5.1),
        Inches(2.9),
        [
            "Төсөл, даалгавар, тайлан олон газар тархчихдаг асуудлыг багасгана.",
            "Хэн юу хийж байгаа, ямар хугацаатай ажиллаж байгааг нэг дороос харуулна.",
            "Supervisor, worker бүтэцтэй учраас байгууллагын дотоод урсгалд илүү тохирно.",
            "AI planner, voice task ашиглаад ажлыг хурдан задлах боломж өгнө.",
        ],
        size=18,
    )
    add_metric(slide, Inches(7.0), Inches(2.9), Inches(2.35), Inches(1.2), "1 платформ", "Төсөл, даалгавар, баг, AI", ACCENT)
    add_metric(slide, Inches(9.55), Inches(2.9), Inches(2.35), Inches(1.2), "3 түвшин", "Admin · Supervisor · Worker", ACCENT_2)
    add_metric(slide, Inches(7.0), Inches(4.35), Inches(4.9), Inches(1.25), "Бодит хэрэглээнд ойр", "Нэвтрэх, эрх, тайлан, company structure, billing, AI зэрэг нь бүгд ажилладаг.", GOOD)


def slide_features(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, RGBColor(245, 247, 251))
    add_title(
        slide,
        "ГОЛ ФУНКЦУУД",
        "Систем юу хийж чаддаг вэ?",
        "Хэрэглэгчийн урсгал, багийн удирдлага, AI туслах, аюулгүй байдлыг нэгтгэсэн.",
    )
    add_panel(slide, Inches(0.8), Inches(2.55), Inches(6.45), Inches(4.2), RGBColor(255, 255, 255), LINE)
    add_bullets(
        slide,
        Inches(1.08),
        Inches(2.9),
        Inches(5.95),
        Inches(3.5),
        [
            "Бүртгэл, нэвтрэлт, logout, forgot password кодоор сэргээх логик.",
            "Role-based access: ADMIN, MODERATOR, USER болон company-level supervisor/worker бүтэц.",
            "Project CRUD, task CRUD, comment, search, filter, pagination.",
            "Dashboard, performance, company workspace, billing, Google login.",
            "AI search, AI planner, voice task capture зэрэг нэмэлт боломжууд.",
        ],
        size=17,
    )
    slide.shapes.add_picture(str(HERO_IMAGE), Inches(7.55), Inches(2.6), Inches(4.35), Inches(4.0))
    add_textbox(slide, Inches(7.72), Inches(6.15), Inches(4.0), Inches(0.4), "Интерфэйс нь responsive бөгөөд desktop, mobile хоёрт ажиллана.", size=12, color=RGBColor(71, 85, 105))


def slide_roles(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, BG)
    add_title(
        slide,
        "ХЭРЭГЛЭГЧИЙН ЛОГИК",
        "Эрхийн бүтэц ба ажлын урсгал",
        "SprintFlow-ийн ялгарах нэг тал нь company hierarchy болон supervisor/worker бүтэц юм.",
        light=True,
    )
    cols = [
        ("System Admin", "Компани үүсгэнэ, хэрэглэгч удирдана, бүх тайлан болон AI боломжийг ашиглана.", ACCENT),
        ("Supervisor", "Өөрийн компанийн ажилтнуудыг багтаа оруулж, төсөл ба даалгавар хуваарилна.", ACCENT_2),
        ("Worker", "Өөрт хамаарах төсөл, даалгавар, inbox болон ажлын явцаа харж шинэчилнэ.", GOOD),
    ]
    left = Inches(0.8)
    for title, desc, color in cols:
        add_panel(slide, left, Inches(2.7), Inches(3.78), Inches(2.75), PANEL, RGBColor(51, 65, 85))
        slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, Inches(2.7), Inches(3.78), Inches(0.12)).fill.solid()
        slide.shapes[-1].fill.fore_color.rgb = color
        slide.shapes[-1].line.fill.background()
        add_textbox(slide, left + Inches(0.22), Inches(3.0), Inches(3.2), Inches(0.4), title, size=20, bold=True, color=TEXT)
        add_textbox(slide, left + Inches(0.22), Inches(3.55), Inches(3.25), Inches(1.45), desc, size=14, color=RGBColor(203, 213, 225))
        left += Inches(4.05)
    add_textbox(slide, Inches(0.85), Inches(6.15), Inches(11.5), Inches(0.35), "Ажлын ерөнхий урсгал: компани үүсгэх → баг холбох → төсөл нээх → даалгавар өгөх → dashboard/performance-ээр хянах", size=13, color=RGBColor(147, 197, 253))


def slide_architecture(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, RGBColor(248, 250, 252))
    add_title(
        slide,
        "АРХИТЕКТУР",
        "Frontend, backend, database нь тусдаа давхаргатай",
        "Энэ бүтэц нь системийг өргөтгөх, deploy хийх, олон хэрэглэгчтэй ажиллуулахад тохиромжтой.",
    )
    labels = [
        ("Browser", "React + Vite frontend"),
        ("Nginx", "Static hosting + /api proxy"),
        ("Express API", "Auth, projects, tasks, dashboard, companies"),
        ("Prisma ORM", "Relational query layer"),
        ("PostgreSQL", "Үндсэн өгөгдөл хадгална"),
    ]
    x = Inches(0.95)
    y = Inches(3.2)
    widths = [1.85, 1.8, 2.35, 1.8, 2.0]
    for idx, ((title, desc), w) in enumerate(zip(labels, widths)):
        add_panel(slide, x, y, Inches(w), Inches(1.55), RGBColor(255, 255, 255), LINE)
        add_textbox(slide, x + Inches(0.16), y + Inches(0.22), Inches(w) - Inches(0.3), Inches(0.35), title, size=18, bold=True, color=ACCENT if idx < 2 else TEXT_DARK)
        add_textbox(slide, x + Inches(0.16), y + Inches(0.63), Inches(w) - Inches(0.3), Inches(0.62), desc, size=12, color=RGBColor(71, 85, 105))
        if idx < len(widths) - 1:
            add_textbox(slide, x + Inches(w) + Inches(0.05), y + Inches(0.55), Inches(0.35), Inches(0.3), "→", size=22, bold=True, color=ACCENT_2, align=PP_ALIGN.CENTER)
        x += Inches(w) + Inches(0.45)
    add_bullets(
        slide,
        Inches(0.95),
        Inches(5.45),
        Inches(11.1),
        Inches(1.3),
        [
            "Database нь users, companies, company_members, projects, project_members, tasks, comments, activity_logs зэрэг гол хүснэгтүүдтэй.",
            "Docker Compose болон Fly.io ашиглаж орчин, deploy-ийг тусгаарласан.",
        ],
        size=15,
        color=TEXT_DARK,
    )


def slide_security(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, RGBColor(245, 247, 251))
    add_title(
        slide,
        "АЮУЛГҮЙ БАЙДАЛ БА ГҮЙЦЭТГЭЛ",
        "Систем зөвхөн харагдах байдал биш, дотоод логикоороо ч хамгаалалттай",
        "Дипломын төсөл боловч production-д ойр шийдлүүдийг сонгосон.",
    )
    add_panel(slide, Inches(0.8), Inches(2.65), Inches(5.6), Inches(3.8), RGBColor(255, 255, 255), LINE)
    add_textbox(slide, Inches(1.05), Inches(2.95), Inches(4.8), Inches(0.35), "Аюулгүй байдлын хэсэг", size=18, bold=True, color=ACCENT)
    add_bullets(
        slide,
        Inches(1.05),
        Inches(3.35),
        Inches(4.95),
        Inches(2.7),
        [
            "bcrypt password hashing, JWT auth, protected routes",
            "CSRF, rate limit, Helmet, HPP, input validation",
            "Company-level data scope ба role-based access control",
            "Forgot password кодоор сэргээх, assignment email мэдэгдэл",
        ],
        size=15,
    )
    add_panel(slide, Inches(6.7), Inches(2.65), Inches(5.55), Inches(3.8), RGBColor(255, 255, 255), LINE)
    add_textbox(slide, Inches(6.95), Inches(2.95), Inches(4.8), Inches(0.35), "Гүйцэтгэл ба өргөтгөх боломж", size=18, bold=True, color=ACCENT_2)
    add_bullets(
        slide,
        Inches(6.95),
        Inches(3.35),
        Inches(4.9),
        Inches(2.7),
        [
            "Search, filter, pagination, indexed relational queries",
            "REST API бүтэц нь олон хэрэглэгч зэрэг холбогдоход илүү тогтвортой",
            "Pricing plan-аар AI entitlement удирдах боломжтой",
            "Цаашид mobile app, real-time collaboration нэмэх суурьтай",
        ],
        size=15,
    )


def slide_conclusion(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, BG)
    add_textbox(slide, Inches(0.8), Inches(0.95), Inches(4.5), Inches(0.4), "ДҮГНЭЛТ", size=12, bold=True, color=RGBColor(147, 197, 253))
    add_textbox(slide, Inches(0.8), Inches(1.35), Inches(6.1), Inches(1.0), "SprintFlow бол багийн ажлыг ойлгомжтой, хяналттай, хурдан болгох зорилготой систем.", size=26, bold=True, color=TEXT)
    add_bullets(
        slide,
        Inches(0.82),
        Inches(2.7),
        Inches(6.0),
        Inches(2.6),
        [
            "Agile төслийн менежментийг нэг платформд төвлөрүүлсэн.",
            "Компанийн бүтэцтэй уялдсан supervisor/worker логиктой.",
            "AI planner, voice task зэрэг нэмэлт боломжтой.",
            "Сургалтын төсөл төдийгүй бодит хэрэглээнд ойр шийдэл болсон.",
        ],
        size=17,
        color=RGBColor(226, 232, 240),
    )
    add_panel(slide, Inches(7.0), Inches(1.55), Inches(4.8), Inches(4.35), PANEL, RGBColor(51, 65, 85))
    add_textbox(slide, Inches(7.35), Inches(1.95), Inches(4.0), Inches(0.45), "Цаашид хөгжүүлэх санаа", size=18, bold=True, color=TEXT)
    add_bullets(
        slide,
        Inches(7.35),
        Inches(2.45),
        Inches(3.95),
        Inches(2.4),
        [
            "Realtime collaboration",
            "Илүү нарийн analytics, BI тайлан",
            "Mobile app эсвэл PWA хувилбар",
            "Enterprise түвшний notification automation",
        ],
        size=15,
        color=RGBColor(203, 213, 225),
    )
    add_textbox(slide, Inches(0.8), Inches(6.55), Inches(5.0), Inches(0.35), "Анхаарал тавьсанд баярлалаа.", size=14, color=TEXT_MUTED)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide_cover(prs)
    slide_problem_goal(prs)
    slide_features(prs)
    slide_roles(prs)
    slide_architecture(prs)
    slide_security(prs)
    slide_conclusion(prs)

    prs.save(OUT)
    print(f"Saved to {OUT}")


if __name__ == "__main__":
    build()
